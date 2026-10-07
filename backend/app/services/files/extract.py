"""Turn uploaded files into text the platform can use.

Documents (PDF/Word/text) become searchable text; spreadsheets become a compact summary
(shape, columns, stats, sample rows) -- the data itself is never executed as code.
"""
import io
from pathlib import Path

DOCUMENT_TYPES = {".pdf", ".docx", ".txt", ".md"}
SPREADSHEET_TYPES = {".csv", ".xlsx"}

_MAX_ROWS = 50_000
_MAX_COLS = 40
_MAX_SAMPLE_ROWS = 20
_MAX_SUMMARY_CHARS = 6000


def extract_document_text(filename: str, data: bytes) -> str:
    suffix = Path(filename).suffix.lower()
    if suffix in (".txt", ".md"):
        return data.decode("utf-8", errors="ignore")
    if suffix == ".pdf":
        from pypdf import PdfReader

        reader = PdfReader(io.BytesIO(data))
        return "\n\n".join(page.extract_text() or "" for page in reader.pages)
    if suffix == ".docx":
        import docx

        document = docx.Document(io.BytesIO(data))
        parts = [p.text for p in document.paragraphs]
        for table in document.tables:
            for row in table.rows:
                parts.append(" | ".join(cell.text for cell in row.cells))
        return "\n".join(parts)
    raise ValueError(f"Unsupported document type: {suffix}")


def summarize_spreadsheet(filename: str, data: bytes) -> str:
    import pandas as pd

    suffix = Path(filename).suffix.lower()
    if suffix == ".csv":
        df = pd.read_csv(io.BytesIO(data), nrows=_MAX_ROWS)
    elif suffix == ".xlsx":
        df = pd.read_excel(io.BytesIO(data), nrows=_MAX_ROWS)
    else:
        raise ValueError(f"Unsupported spreadsheet type: {suffix}")

    cols = list(df.columns)[:_MAX_COLS]
    summary = (
        f"Rows: {len(df)}{' (only the first %d were read)' % _MAX_ROWS if len(df) == _MAX_ROWS else ''}\n"
        f"Columns ({len(df.columns)}): {[str(c) for c in cols]}\n\n"
        f"Column types: {df[cols].dtypes.astype(str).to_dict()}\n\n"
        f"Summary statistics:\n{df[cols].describe(include='all').to_string()}\n\n"
        f"First rows:\n{df[cols].head(_MAX_SAMPLE_ROWS).to_string()}"
    )
    return summary[:_MAX_SUMMARY_CHARS]
